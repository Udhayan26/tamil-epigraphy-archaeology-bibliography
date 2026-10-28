#!/usr/bin/env python3
"""Write data/Tamil_Epigraphy_Books.xlsx from data/bibliography.json (+ heritage records)."""
import json, re, sys
from datetime import date, timedelta
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = Path(__file__).resolve().parent.parent
DATA = ROOT / "data"
OUT = DATA / "Tamil_Epigraphy_Books.xlsx"

FONT = "Arial"
HEAD_FILL = PatternFill("solid", fgColor="7A3A12")
HEAD_FONT = Font(name=FONT, bold=True, color="FFFFFF", size=10)
BODY = Font(name=FONT, size=10)
LINK = Font(name=FONT, size=10, color="0563C1", underline="single")
TITLE = Font(name=FONT, size=16, bold=True, color="7A3A12")
THIN = Side(style="thin", color="D9CFC2")

SHEET = {"EPG": "Epigraphy Studies", "ASI": "SII, EI & Colonial Series", "TNS": "Tamil Nadu Inscriptions",
         "CPP": "Copper Plates", "TBR": "Tamil-Brahmi & Scripts", "HER": "Hero Stones", "ROC": "Rock Art",
         "EXC": "Excavation Reports", "PRE": "Prehistory & Megaliths", "ARC": "Archaeology General",
         "HIS": "History from Inscriptions", "NUM": "Coins", "PLM": "Palm-leaf Manuscripts", "SEA": "Seals & Rock-cut Caves",
         "REF": "Catalogues & Glossaries", "JOU": "Journals", "ONL": "Online Resources"}
ACCESS = {"open": "Free to read", "open-us": "Free (USA only)", "preview": "Preview", "borrow": "Borrow",
          "catalogue": "Not online – buy / library", "": "Not online – buy / library"}

COLS = [  # header, width, getter
    ("ID", 11, lambda r: r["id"]),
    ("Category", 30, lambda r: r["category_name"]),
    ("வகை", 26, lambda r: r["category_name_ta"]),
    ("Title (Tamil)", 42, lambda r: r["title_ta"]),
    ("Title (English / romanised)", 48, lambda r: r["title_en"] or r["title"]),
    ("Author / Editor", 30, lambda r: r["authors"] + (f" ({r['authors_roman']})" if r.get("authors_roman") else "")),
    ("Publisher", 32, lambda r: r["publisher"] + (f" ({r['publisher_roman']})" if r.get("publisher_roman") else "")),
    ("Place", 14, lambda r: r["place"]),
    ("Year", 8, lambda r: r["year"]),
    ("Language", 12, lambda r: r["language"]),
    ("Type", 13, lambda r: r["type"]),
    ("Series / Volume", 18, lambda r: r.get("series", "")),
    ("ISBN", 15, lambda r: r.get("isbn", "")),
    ("Access", 22, lambda r: ACCESS.get(r["access"], r["access"])),
    ("Read online", 12, "read"),
    ("Download PDF", 13, "pdf"),
    ("More copies", 12, "copies"),
    ("Buy", 10, "buy"),
    ("Source record", 13, "source"),
    ("Other records", 13, "others"),
    ("Date added", 11, lambda r: date.fromisoformat(r["added"]) if re.match(r"\d{4}-\d\d-\d\d$", r.get("added", "")) else None),
    ("Notes", 40, lambda r: r.get("notes", "")),
]


def buy_url(r):
    if r.get("buy_url"):
        return r["buy_url"]
    if r["access"] in ("open", "open-us"):
        return ""
    from urllib.parse import quote_plus
    if r.get("title_ta"):
        return "https://www.panuval.com/index.php?route=product/search&search=" + quote_plus(r["title_ta"][:80])
    q = ((r.get("title_en") or r["title"])[:100] + " " + (r["authors"].split(";")[0] if r["authors"] else "")).strip()
    return "https://www.amazon.in/s?k=" + quote_plus(q)


def link(ws, row, col, url, text):
    if not url:
        return
    c = ws.cell(row=row, column=col, value=text)
    c.hyperlink = url[:2000]
    c.font = LINK


def book_sheet(wb, name, rows, table_name):
    ws = wb.create_sheet(name)
    for i, (h, w, _) in enumerate(COLS, 1):
        c = ws.cell(row=1, column=i, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
        c.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    ws.row_dimensions[1].height = 30
    for ri, r in enumerate(rows, 2):
        for ci, (h, _, g) in enumerate(COLS, 1):
            if callable(g):
                v = g(r)
                c = ws.cell(row=ri, column=ci, value=v if v not in (None, "") else None)
                c.font = BODY
                c.alignment = Alignment(vertical="top", wrap_text=ci in (4, 5, 6, 7, 22))
                if h == "Date added":
                    c.number_format = "yyyy-mm-dd"
            elif g == "read" and r["access"] in ("open", "open-us"):
                link(ws, ri, ci, r.get("read_url"), "Read")
            elif g == "pdf":
                link(ws, ri, ci, r.get("download_url"), "PDF")
            elif g == "copies" and r.get("copies"):
                c0 = r["copies"][0]
                link(ws, ri, ci, c0.get("pdf") or c0.get("read"), f"{len(r['copies'])} more")
            elif g == "buy":
                link(ws, ri, ci, buy_url(r), "Buy")
            elif g == "source":
                link(ws, ri, ci, r.get("source_url"), "Record")
            elif g == "others" and r.get("other_source_urls"):
                link(ws, ri, ci, r["other_source_urls"][0], f"{len(r['other_source_urls'])} more")
    last = max(2, len(rows) + 1)
    ref = f"A1:{get_column_letter(len(COLS))}{last}"
    t = Table(displayName=table_name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)
    ws.freeze_panes = "F2"
    return ws


def main():
    doc = json.loads((DATA / "bibliography.json").read_text(encoding="utf-8"))
    R, cats = doc["records"], doc["categories"]
    today = date.today()
    wb = Workbook()

    # ---- Summary
    ws = wb.active
    ws.title = "Summary"
    ws["A1"] = "Tamil Epigraphy, Inscriptions, Archaeology, Coins, Copper Plates, Palm-leaf & Hero Stones — Books"
    ws["A1"].font = TITLE
    ws["A2"] = "தமிழ்க் கல்வெட்டியல், கல்வெட்டுகள், தொல்லியல், நாணயங்கள், செப்பேடுகள், ஓலைச்சுவடிகள், நடுகற்கள் — நூல் பட்டியல்"
    ws["A2"].font = Font(name=FONT, size=12, color="7A3A12")
    ws["A3"] = f"Updated {doc['generated']}. One row per book or volume (duplicates merged). Click Read / PDF / Buy / Record to open links."
    ws["A3"].font = Font(name=FONT, size=10, italic=True, color="6F655C")
    n = len(R) + 1
    B = "'All Books'"
    stats = [
        ("Total books & volumes", f"=COUNTA({B}!A2:A{n})"),
        ("With author / editor", f"=COUNTIF({B}!F2:F{n},\"?*\")"),
        ("With publisher", f"=COUNTIF({B}!G2:G{n},\"?*\")"),
        ("With year", f"=COUNTIF({B}!I2:I{n},\"?*\")"),
        ("Free to read online", f"=COUNTIF({B}!N2:N{n},\"Free*\")"),
        ("With direct PDF download", f"=COUNTIF({B}!P2:P{n},\"PDF\")"),
        ("Not online – buy / library", f"=COUNTIF({B}!N2:N{n},\"Not online*\")"),
        ("Tamil-language", f"=COUNTIF({B}!J2:J{n},\"Tamil*\")"),
        ("Added in last 7 days", f"=COUNTIF({B}!U2:U{n},\">=\"&(TODAY()-7))"),
        ("Added in last 30 days", f"=COUNTIF({B}!U2:U{n},\">=\"&(TODAY()-30))"),
    ]
    ws["A5"], ws["B5"] = "Overview", "Count"
    for c in (ws["A5"], ws["B5"]):
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    for i, (label, f) in enumerate(stats, 6):
        ws.cell(row=i, column=1, value=label).font = BODY
        c = ws.cell(row=i, column=2, value=f)
        c.font, c.number_format = BODY, "#,##0"
    r0 = 6 + len(stats) + 1
    for j, h in enumerate(["Category", "வகை", "Books", "Free to read", "Sheet"], 1):
        c = ws.cell(row=r0, column=j, value=h)
        c.font, c.fill = HEAD_FONT, HEAD_FILL
    sheet_names = {}
    for i, (code, v) in enumerate(cats.items(), r0 + 1):
        nm = SHEET.get(code) or re.sub(r"[\\/*?:\[\]]", "", v["en"])[:28]
        sheet_names[code] = nm
        ws.cell(row=i, column=1, value=v["en"]).font = BODY
        ws.cell(row=i, column=2, value=v["ta"]).font = BODY
        c = ws.cell(row=i, column=3, value=f"=COUNTIF({B}!B2:B{n},A{i})")
        c.font, c.number_format = BODY, "#,##0"
        c = ws.cell(row=i, column=4, value=f"=COUNTIFS({B}!B2:B{n},A{i},{B}!N2:N{n},\"Free*\")")
        c.font, c.number_format = BODY, "#,##0"
        lc = ws.cell(row=i, column=5, value=nm)
        lc.hyperlink = f"#'{nm}'!A1"
        lc.font = LINK
    endr = r0 + len(cats)
    ws.cell(row=endr + 1, column=1, value="Total").font = Font(name=FONT, bold=True, size=10)
    c = ws.cell(row=endr + 1, column=3, value=f"=SUM(C{r0 + 1}:C{endr})")
    c.font, c.number_format = Font(name=FONT, bold=True, size=10), "#,##0"
    note = endr + 3
    for k, txt in enumerate([
        "How to use: open 'All Books' and use the filter arrows (Category, Language, Year, Access). Each category also has its own sheet.",
        "Access: 'Free to read' = full book online (Read / PDF links). 'Not online' = use Buy (Tamil books: Panuval search; others: Amazon.in search) or Record to find a library copy.",
        "More copies = other scans of the same book or volume. Other records = the same book in other catalogues.",
        "Titles: Tamil script and English / exact romanised form (ISO 15919) where the catalogue gave only one script.",
        "Details come from library catalogues; please check against the book itself.",
    ]):
        ws.cell(row=note + k, column=1, value=txt).font = Font(name=FONT, size=9, color="6F655C")
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 40
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 30

    # ---- data sheets
    book_sheet(wb, "All Books", R, "AllBooks")
    cutoff = (today - timedelta(days=30)).isoformat()
    new = sorted([r for r in R if r.get("added", "") >= cutoff and not r.get("curated")], key=lambda r: r.get("added", ""), reverse=True)
    book_sheet(wb, "Newly Added (30 days)", new, "NewlyAdded")
    for i, (code, v) in enumerate(cats.items()):
        rs = [r for r in R if r["category"] == code]
        if rs:
            book_sheet(wb, sheet_names[code], rs, f"Cat_{code}")

    # ---- photo / artefact records
    hp = DATA / "heritage_records.json"
    if hp.exists():
        H = json.loads(hp.read_text(encoding="utf-8"))
        ws = wb.create_sheet("Photo Records (not books)")
        for i, (h, w) in enumerate([("Title (Tamil)", 40), ("Title (romanised)", 40), ("Category", 10), ("Record", 12)], 1):
            c = ws.cell(row=1, column=i, value=h)
            c.font, c.fill = HEAD_FONT, HEAD_FILL
            ws.column_dimensions[get_column_letter(i)].width = w
        for ri, h in enumerate(H, 2):
            ws.cell(row=ri, column=1, value=h["title_ta"]).font = BODY
            ws.cell(row=ri, column=2, value=h["title_en"]).font = BODY
            ws.cell(row=ri, column=3, value=h["category"]).font = BODY
            link(ws, ri, 4, h["url"], "Record")
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:D{max(2, len(H) + 1)}"

    wb.calculation.fullCalcOnLoad = True
    wb.save(OUT)
    print(f"Excel written: {OUT.name} ({len(R)} books, {len(wb.sheetnames)} sheets)")


if __name__ == "__main__":
    main()
